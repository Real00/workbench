from io import BytesIO

from progress.domain import decode_unicode_text

SUPPORTED_IMPORTS = {".md", ".txt", ".docx"}


def extract_markdown(filename: str, data: bytes) -> str:
    name = filename.replace("\\", "/").rsplit("/", 1)[-1].strip()
    if not name:
        raise ValueError("filename is required")
    suffix = f".{name.rsplit('.', 1)[-1].lower()}" if "." in name else ""
    if suffix not in SUPPORTED_IMPORTS:
        raise ValueError("unsupported document type")
    if suffix in {".md", ".txt"}:
        try:
            text = data.decode("utf-8")
        except UnicodeDecodeError as extra:
            raise ValueError("markdown must be UTF-8") from extra
        return decode_unicode_text(text.replace("\r\n", "\n").replace("\r", "\n")).strip()
    return _docx_to_markdown(data)


def _docx_to_markdown(data: bytes) -> str:
    from docx import Document

    try:
        document = Document(BytesIO(data))
    except Exception as extra:
        raise ValueError("could not read Word document") from extra
    blocks: list[str] = []
    for paragraph in document.paragraphs:
        text = decode_unicode_text(paragraph.text).strip()
        if text:
            blocks.append(text)
    for table in document.tables:
        rendered = _table_markdown(table)
        if rendered:
            blocks.append(rendered)
    body = "\n\n".join(blocks).strip()
    if not body:
        raise ValueError("Word document has no extractable text")
    return body


def _table_markdown(table: object) -> str:
    rows: list[list[str]] = []
    for row in getattr(table, "rows", []):
        cells = [decode_unicode_text(cell.text).replace("\n", " ").strip() for cell in row.cells]
        if any(cells):
            rows.append(cells)
    if not rows:
        return ""
    width = max(len(row) for row in rows)
    padded = [row + [""] * (width - len(row)) for row in rows]
    header = padded[0]
    lines = [
        "| " + " | ".join(header) + " |",
        "| " + " | ".join("---" for _ in header) + " |",
    ]
    for row in padded[1:]:
        lines.append("| " + " | ".join(row) + " |")
    return "\n".join(lines)
